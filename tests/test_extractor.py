from io import BytesIO

import fitz
import httpx
import pytest
from docx import Document

from talent_policy_search.config import Settings
from talent_policy_search.extractor import extract_fulltext, extract_page_info
from talent_policy_search.fetcher import FetchedPage, Fetcher, FetchTooLargeError


class ChunkedAsyncStream(httpx.AsyncByteStream):
    def __init__(self, chunks: list[bytes]) -> None:
        self._chunks = chunks

    async def __aiter__(self):
        for chunk in self._chunks:
            yield chunk


class UnreadableAsyncStream(httpx.AsyncByteStream):
    async def __aiter__(self):
        raise AssertionError("response body should not be read")
        yield b""


def test_extract_html_fulltext_removes_script_and_keeps_body_text():
    html = """
    <html>
      <head><title>Example</title><script>bad()</script></head>
      <body><h1>人才政策</h1><p>给予高层次人才住房补贴。</p></body>
    </html>
    """.encode()

    text = extract_fulltext(html, "text/html", max_chars=200)

    assert "人才政策" in text
    assert "住房补贴" in text
    assert "bad()" not in text


def test_extract_page_info_extracts_policy_items_dates_and_action_links():
    html = """
    <html>
      <head><title>人才政策-广州市人力资源和社会保障局网站</title></head>
      <body>
        <h1>人才政策</h1>
        <ul>
          <li>
            <a href="/ywzt/rcgz/renczc/content/post_1.html">关于做好广州市境外人才财政补贴申报准备工作的通知</a>
            <span>[ 2025-03-12 ]</span>
          </li>
          <li>
            <a href="https://evil.example/bad.html">非官方转载</a>
            <span>2025-04-01</span>
          </li>
        </ul>
        <a href="/apply/index.html">在线申报入口</a>
        <a href="/files/material.docx">申报材料下载</a>
      </body>
    </html>
    """.encode()

    info = extract_page_info(
        html,
        "text/html; charset=utf-8",
        max_chars=1000,
        base_url="https://rsj.gz.gov.cn/ywzt/rcgz/renczc/",
    )

    assert info.title == "人才政策"
    assert "关于做好广州市境外人才财政补贴申报准备工作的通知" in info.fulltext
    assert info.policy_items[0].title == "关于做好广州市境外人才财政补贴申报准备工作的通知"
    assert info.policy_items[0].url == "https://rsj.gz.gov.cn/ywzt/rcgz/renczc/content/post_1.html"
    assert info.policy_items[0].published_date == "2025-03-12"
    assert info.policy_items[0].snippet
    assert info.application_links[0].text == "在线申报入口"
    assert info.attachment_links[0].url == "https://rsj.gz.gov.cn/files/material.docx"


def test_extract_html_filters_navigation_links_from_policy_items():
    html = """
    <html>
      <body>
        <a href="/nav/">学校首页</a>
        <a href="/rsc/office">人事科</a>
        <a href="/policy/2025/notice.html">2025年人才发展专项支持计划（试点）</a>
        <a href="/policy/2025/download.zip">资料下载</a>
      </body>
    </html>
    """.encode()

    info = extract_page_info(
        html,
        "text/html",
        max_chars=800,
        base_url="https://www.example.gov.cn/",
    )

    assert len(info.policy_items) == 1
    assert info.policy_items[0].title == "2025年人才发展专项支持计划（试点）"


def test_extract_html_filters_policy_library_links_without_talent_signal():
    html = """
    <html>
      <body>
        <a href="/nav/file-library">政策文件库</a>
        <a href="/policy/2025-notice.html">2025年高层次人才专项申报通知</a>
      </body>
    </html>
    """.encode()

    info = extract_page_info(
        html,
        "text/html",
        max_chars=800,
        base_url="https://www.example.gov.cn/",
    )

    assert len(info.policy_items) == 1
    assert info.policy_items[0].title == "2025年高层次人才专项申报通知"


def test_extract_page_info_uses_document_title_for_detail_pages_without_policy_list():
    html = """
    <html>
      <head><title>2025年黄埔区人才津贴申报通知</title></head>
      <body>
        <h1>2025年黄埔区人才津贴申报通知</h1>
        <p>人才津贴发放标准为每年最高20万元，补贴期3年。</p>
      </body>
    </html>
    """.encode()

    info = extract_page_info(
        html,
        "text/html",
        max_chars=1000,
        base_url="http://www.hp.gov.cn/policy.html",
    )

    assert info.title == "2025年黄埔区人才津贴申报通知"
    assert info.policy_items == []
    assert info.date_candidates == []


def test_extract_html_fulltext_tolerates_charset_parameter():
    html = "<html><body><p>深圳人才奖励</p></body></html>".encode()

    text = extract_fulltext(html, "text/html; charset=utf-8", max_chars=200)

    assert text == "深圳人才奖励"


def test_extract_html_fulltext_uses_charset_parameter_for_decoding():
    html = "<html><body><p>深圳人才奖励</p></body></html>".encode("gbk")

    text = extract_fulltext(html, "text/html; charset=gbk", max_chars=200)

    assert text == "深圳人才奖励"


def test_extract_html_fulltext_respects_character_limit():
    html = b"<html><body><p>abcdefgh</p></body></html>"

    text = extract_fulltext(html, "text/html", max_chars=4)

    assert text == "abcd"
    assert len(text) == 4


def test_extract_plain_text_respects_character_limit():
    text = extract_fulltext("abcdef".encode(), "text/plain", max_chars=3)

    assert text == "abc"


def test_extract_plain_text_uses_charset_parameter_for_decoding():
    text = extract_fulltext(
        "深圳人才奖励".encode("gbk"),
        "text/plain; charset=gbk",
        max_chars=200,
    )

    assert text == "深圳人才奖励"


def test_extract_pdf_fulltext_uses_in_memory_bytes():
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "Talent policy housing subsidy")
    content = document.tobytes()
    document.close()

    text = extract_fulltext(content, "application/pdf", max_chars=200)

    assert "Talent policy housing subsidy" in text


def test_extract_invalid_pdf_returns_empty_string():
    text = extract_fulltext(b"not a pdf", "application/pdf", max_chars=200)

    assert text == ""


def test_extract_pdf_stops_reading_pages_after_max_chars(monkeypatch):
    calls = []

    class FakePage:
        def __init__(self, text: str, page_number: int) -> None:
            self._text = text
            self._page_number = page_number

        def get_text(self, kind: str) -> str:
            assert kind == "text"
            calls.append(self._page_number)
            return self._text

    class FakeDocument:
        def __enter__(self):
            return self

        def __exit__(self, exc_type, exc, traceback):
            return None

        def __iter__(self):
            return iter(
                [
                    FakePage("first page has enough characters", 1),
                    FakePage("later page should not be read", 2),
                ]
            )

    def fake_open(*, stream: bytes, filetype: str):
        assert stream == b"pdf bytes"
        assert filetype == "pdf"
        return FakeDocument()

    monkeypatch.setattr("talent_policy_search.extractor.fitz.open", fake_open)

    text = extract_fulltext(b"pdf bytes", "application/pdf", max_chars=10)

    assert text == "first page"
    assert calls == [1]


def test_extract_docx_fulltext_uses_in_memory_bytes():
    document = Document()
    document.add_paragraph("高层次人才奖励")
    buffer = BytesIO()
    document.save(buffer)

    text = extract_fulltext(
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        max_chars=200,
    )

    assert text == "高层次人才奖励"


def test_extract_invalid_docx_returns_empty_string():
    text = extract_fulltext(
        b"not a docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        max_chars=200,
    )

    assert text == ""


def test_extract_docx_stops_reading_paragraphs_after_max_chars(monkeypatch):
    calls = []

    class FakeParagraph:
        def __init__(self, text: str, paragraph_number: int) -> None:
            self._text = text
            self._paragraph_number = paragraph_number

        @property
        def text(self) -> str:
            calls.append(self._paragraph_number)
            return self._text

    class FakeDocument:
        paragraphs = [
            FakeParagraph("first paragraph has enough characters", 1),
            FakeParagraph("later paragraph should not be read", 2),
        ]

    def fake_document(content: BytesIO):
        assert content.getvalue() == b"docx bytes"
        return FakeDocument()

    monkeypatch.setattr("talent_policy_search.extractor.Document", fake_document)

    text = extract_fulltext(
        b"docx bytes",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        max_chars=15,
    )

    assert text == "first paragraph"
    assert calls == [1]


def test_extract_unsupported_content_type_returns_empty_string():
    text = extract_fulltext(b"{}", "application/json", max_chars=200)

    assert text == ""


class FakeScraplingFetcher:
    def __init__(self) -> None:
        self.urls: list[str] = []

    async def fetch(self, url: str) -> FetchedPage:
        self.urls.append(url)
        return FetchedPage(
            url=url,
            final_url=url,
            content_type="text/html; charset=utf-8",
            content=b"<html><body>scrapling policy</body></html>",
        )


@pytest.mark.asyncio
async def test_fetcher_sends_retrieval_headers_and_returns_page_content():
    async def handler(request: httpx.Request) -> httpx.Response:
        assert request.headers["user-agent"] == (
            "TalentPolicyLiveSearch/0.1 (+official-source-retrieval)"
        )
        assert request.headers["accept"] == (
            "text/html,application/xhtml+xml,application/pdf,text/plain,*/*;q=0.8"
        )
        return httpx.Response(
            200,
            content=b"<html><body>policy</body></html>",
            headers={"content-type": "text/html; charset=utf-8"},
            request=request,
        )

    fetcher = Fetcher(
        settings=Settings(SEARCH_TIMEOUT_SECONDS=45, _env_file=None),
        transport=httpx.MockTransport(handler),
    )

    page = await fetcher.fetch("https://example.gov.cn/policy")

    assert page.url == "https://example.gov.cn/policy"
    assert page.final_url == "https://example.gov.cn/policy"
    assert page.content_type == "text/html; charset=utf-8"
    assert page.content == b"<html><body>policy</body></html>"


@pytest.mark.asyncio
async def test_fetcher_uses_scrapling_fallback_for_bad_ecpoint_when_enabled():
    async def handler(request: httpx.Request) -> httpx.Response:
        raise httpx.ConnectError("[SSL: BAD_ECPOINT] bad ecpoint", request=request)

    scrapling_fetcher = FakeScraplingFetcher()
    fetcher = Fetcher(
        settings=Settings(ENABLE_SCRAPLING_FALLBACK=True, _env_file=None),
        transport=httpx.MockTransport(handler),
        scrapling_fetcher=scrapling_fetcher,
    )

    page = await fetcher.fetch("https://www.sz.gov.cn/")

    assert page.content == b"<html><body>scrapling policy</body></html>"
    assert scrapling_fetcher.urls == ["https://www.sz.gov.cn/"]


@pytest.mark.asyncio
async def test_fetcher_uses_scrapling_fallback_for_generic_http_error_when_enabled():
    async def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(404, content=b"not found", request=request)

    scrapling_fetcher = FakeScraplingFetcher()
    fetcher = Fetcher(
        settings=Settings(ENABLE_SCRAPLING_FALLBACK=True, _env_file=None),
        transport=httpx.MockTransport(handler),
        scrapling_fetcher=scrapling_fetcher,
    )

    page = await fetcher.fetch("https://example.gov.cn/policy")

    assert page.content == b"<html><body>scrapling policy</body></html>"
    assert scrapling_fetcher.urls == ["https://example.gov.cn/policy"]


@pytest.mark.asyncio
async def test_fetcher_uses_legacy_tls_context_for_handshake_failure():
    calls: list[str] = []

    async def handler(request: httpx.Request) -> httpx.Response:
        calls.append(str(request.url))
        if len(calls) == 1:
            raise httpx.ConnectError(
                "[SSL: SSLV3_ALERT_HANDSHAKE_FAILURE] ssl/tls alert handshake failure",
                request=request,
            )
        return httpx.Response(
            200,
            content=b"<html><body>legacy handshake policy</body></html>",
            headers={"content-type": "text/html; charset=utf-8"},
            request=request,
        )

    fetcher = Fetcher(
        settings=Settings(SEARCH_TIMEOUT_SECONDS=45, _env_file=None),
        transport=httpx.MockTransport(handler),
    )

    page = await fetcher.fetch("https://example.gov.cn/policy")

    assert page.content == b"<html><body>legacy handshake policy</body></html>"
    assert calls == [
        "https://example.gov.cn/policy",
        "https://example.gov.cn/policy",
    ]


@pytest.mark.asyncio
async def test_fetcher_follows_redirects_and_uses_content_type_fallback():
    async def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path == "/start":
            return httpx.Response(
                302,
                headers={"location": "https://example.gov.cn/final"},
                request=request,
            )
        return httpx.Response(200, content=b"final", request=request)

    fetcher = Fetcher(
        settings=Settings(SEARCH_TIMEOUT_SECONDS=45, _env_file=None),
        transport=httpx.MockTransport(handler),
    )

    page = await fetcher.fetch("https://example.gov.cn/start")

    assert page.url == "https://example.gov.cn/start"
    assert page.final_url == "https://example.gov.cn/final"
    assert page.content_type == "application/octet-stream"
    assert page.content == b"final"


@pytest.mark.asyncio
async def test_fetcher_rejects_disallowed_redirect_without_fetching_target():
    requested_urls: list[str] = []

    async def handler(request: httpx.Request) -> httpx.Response:
        requested_urls.append(str(request.url))
        if str(request.url) == "https://hrss.sz.gov.cn/policy.html":
            return httpx.Response(
                302,
                headers={"location": "https://evil.example/secret"},
                request=request,
            )
        raise AssertionError(f"unexpected request: {request.url}")

    fetcher = Fetcher(
        settings=Settings(SEARCH_TIMEOUT_SECONDS=45, _env_file=None),
        transport=httpx.MockTransport(handler),
    )

    with pytest.raises(RuntimeError, match="Disallowed redirect"):
        await fetcher.fetch(
            "https://hrss.sz.gov.cn/policy.html",
            allowed_redirect_url=lambda url: url.endswith(".sz.gov.cn/policy.html"),
        )

    assert requested_urls == ["https://hrss.sz.gov.cn/policy.html"]


@pytest.mark.asyncio
async def test_fetcher_raises_for_http_error_status():
    async def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(404, content=b"missing", request=request)

    fetcher = Fetcher(
        settings=Settings(
            SEARCH_TIMEOUT_SECONDS=45,
            ENABLE_SCRAPLING_FALLBACK=False,
            _env_file=None,
        ),
        transport=httpx.MockTransport(handler),
    )

    with pytest.raises(httpx.HTTPStatusError):
        await fetcher.fetch("https://example.gov.cn/missing")


@pytest.mark.asyncio
async def test_fetcher_rejects_content_length_over_byte_limit_before_reading():
    async def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            headers={"content-length": "6"},
            stream=UnreadableAsyncStream(),
            request=request,
        )

    fetcher = Fetcher(
        settings=Settings(MAX_FETCH_BYTES_PER_PAGE=5, _env_file=None),
        transport=httpx.MockTransport(handler),
    )

    with pytest.raises(FetchTooLargeError, match="exceeds max fetch bytes"):
        await fetcher.fetch("https://example.gov.cn/too-large")


@pytest.mark.asyncio
async def test_fetcher_rejects_streamed_body_over_byte_limit():
    async def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            stream=ChunkedAsyncStream([b"abc", b"def"]),
            request=request,
        )

    fetcher = Fetcher(
        settings=Settings(MAX_FETCH_BYTES_PER_PAGE=5, _env_file=None),
        transport=httpx.MockTransport(handler),
    )

    with pytest.raises(FetchTooLargeError, match="exceeds max fetch bytes"):
        await fetcher.fetch("https://example.gov.cn/too-large")
